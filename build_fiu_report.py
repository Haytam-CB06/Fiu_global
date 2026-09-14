from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Fiu_Global_Portal_Internship_Report_Humanized_Final.docx"
LOGO = ROOT / "wwwroot" / "img" / "fiu9-mark2.png"

NAVY = "183B66"
BLUE = "2F6FB5"
LIGHT_BLUE = "EAF2FA"
PALE_BLUE = "F5F9FD"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "6B7280"
BLACK = "000000"

TOP_LEVEL_NUMBERS = {
    "Executive summary": "1",
    "Project background and objectives": "2",
    "Portal users and scope": "3",
    "System design": "4",
    "User portal features": "5",
    "Administration features": "6",
    "Security and privacy": "7",
    "Accessibility and language support": "8",
    "Internship work and learning": "9",
    "Internship reflection": "9.1",
    "Verification and results": "10",
    "Constraints and risks": "11",
    "Recommended next work": "12",
    "Delivery assessment": "13",
    "How to use the FIU Global Portal": "14",
    "Portal screenshots, sign in and dashboard": "14.1",
    "Portal screenshots, dining": "14.2",
    "Portal screenshots, chat and administration": "14.3",
    "Extended project record": "15",
    "Appendix Route summary": "A"
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_GRAY, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_run_font(run, name="Aptos", size=10.5, color=BLACK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(paragraph, before=0, after=7, line=1.12, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if keep:
        fmt.keep_with_next = True


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MID_GRAY)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    display_text = text
    if level == 1 and text in TOP_LEVEL_NUMBERS:
        display_text = f"{TOP_LEVEL_NUMBERS[text]} {text}"
    run = paragraph.add_run(display_text)
    size = 15 if level == 1 else 12
    set_run_font(run, size=size, color=BLACK, bold=True)
    style_paragraph(paragraph, before=14 if level == 1 else 9, after=5, line=1.0, keep=True)
    return paragraph


HUMANIZED_REWRITES = [
    ("This report presents the FIU Global Portal, an internship project built to give students, instructors, and administrators one reliable entry point for university services. The portal brings service links, announcements, dining information, academic directory records, notifications, account settings, and campus chat into one sign-in flow.", "I prepared this report to record my internship work on the FIU Global Portal. The portal puts service links, announcements, dining information, directory records, notifications, account settings, and campus chat behind one sign-in flow. My focus was to explain how the project works and what still needs testing before a live release."),
    ("I reviewed the source structure, followed the main request paths, checked the local runtime, and recorded the build result. The project builds with zero warnings and zero errors. Local requests show the expected login redirect, a working login page, and protected data routes for users without a session. The next release needs automated journey tests, production monitoring, and a documented backup process.", "I read the source structure, traced the main request paths, ran the local host, and recorded the build result. The build returned zero warnings and zero errors. I also checked the login redirect and protected routes without a session. The next release needs automated journey tests, production monitoring, and a documented backup process."),
    ("The project gave me practice across the full web application path. I worked with C# route handlers, MySQL state storage, browser-side JavaScript, responsive CSS, form validation, file import logic, and access rules. I also followed how a login request becomes a session, how a role changes the visible dashboard, and how a user action reaches the database.", "During the internship, I worked across the full web application path. My tasks covered C# route handlers, MySQL state storage, browser-side JavaScript, responsive CSS, form validation, file import logic, and access rules. I followed how a login request becomes a session, how a role changes the dashboard, and how a user action reaches the database."),
    ("University services often run on separate sites. Users must remember several links and check each service for new information. This project addresses the access problem with one portal shell and a role-aware dashboard. A signed-in user sees assigned services, then opens the required destination from one screen.", "I started from a simple problem: university services sat on separate sites, so users had to remember several links. The portal brings those entry points into one role-aware dashboard. After sign-in, a user sees assigned services and opens the needed destination from one screen."),
    ("The server uses minimal API routes in Program.cs. AppDataStore contains the main portal operations. DatabaseStateStorage keeps the full state document in the dotnet_app_state table and mirrors core records into relational tables where older tools still need them. Sessions use dotnet_sessions. Chat messages use dotnet_chat_messages. Faculties and departments use linked directory tables.", "I traced the server through minimal API routes in Program.cs. AppDataStore handles the main portal operations. DatabaseStateStorage keeps the broad state document in the dotnet_app_state table and mirrors core records into relational tables for older tools. Sessions use dotnet_sessions. Chat messages use dotnet_chat_messages. Faculties and departments use linked directory tables."),
    ("The project objectives were practical. Give users one trusted starting point. Show daily campus information in a short dashboard view. Give administrators direct control over content and access. Keep older accounts available during the migration from the earlier PHP setup. Protect account and chat data during normal use.", "I set practical project objectives. I wanted one trusted starting point for users, a short dashboard view for daily information, direct content controls for administrators, continued access for older accounts during the PHP migration, and protection for account and chat data."),
    ("The application creates missing tables during startup.", "At startup, the application creates missing tables."),
    ("This report presents the FIU Global Portal", "I prepared this report to document the FIU Global Portal"),
    ("I reviewed the source structure, followed", "I read the source structure and traced"),
    ("The report keeps starter records", "I kept starter records"),
    ("University services often run on separate sites.", "At the start of this review, university services ran on separate sites."),
    ("Students need fast access", "I wrote the scope around daily needs. Students need fast access"),
    ("The user portal covers", "I mapped the user portal to"),
    ("The portal links to", "I recorded links to"),
    ("The server uses minimal API routes", "I traced the server through minimal API routes"),
    ("The browser client loads page shells", "In the browser, the client loads page shells"),
    ("The login page accepts local credentials", "I checked the login page with local credentials"),
    ("Platform cards group services", "I saw platform cards group services"),
    ("The dining page uses a month calendar.", "I checked the dining page, which uses a month calendar."),
    ("The profile page supports", "I checked profile updates for"),
    ("The administrator dashboard presents", "I reviewed the administrator dashboard, which presents"),
    ("Account management covers", "I reviewed account management for"),
    ("The academic directory stores", "I traced the academic directory, which stores"),
    ("Passwords use salted PBKDF2", "I confirmed the use of salted PBKDF2"),
    ("Session cookies use HttpOnly", "I checked HttpOnly"),
    ("Chat text uses AES-GCM", "I confirmed AES-GCM"),
    ("Role checks protect", "I checked role checks on"),
    ("The pages include labels", "I checked the pages for labels"),
    ("The portal includes English", "I confirmed support for English"),
    ("The project gave me practice", "During the internship, I worked"),
    ("The portal made security work concrete.", "Security work became concrete during the review."),
    ("The main lesson was the value", "My main lesson was the value"),
    ("I ran the project build", "I ran the project build"),
    ("The repository has no dedicated", "I found no dedicated"),
    ("Production use needs", "A live release needs"),
    ("FIU Global Portal has a broad working base", "My review found a broad working base in FIU Global Portal"),
    ("The build and basic runtime checks pass.", "The build and basic runtime checks passed during my review."),
    ("The portal serves several groups", "I saw several groups using the portal"),
    ("The map also shows", "The map showed"),
    ("The functional requirements describe", "I grouped the functional requirements around"),
    ("The portal also has quality requirements.", "I treated the following quality points as release checks."),
    ("The domain model holds", "I traced the domain model and found"),
    ("DatabaseStateStorage keeps", "I traced DatabaseStateStorage, which keeps"),
    ("Program.cs holds", "I traced Program.cs, which holds"),
    ("The local sign-in flow starts", "I followed the local sign-in flow"),
    ("The portal also accepts Google sign-in", "I reviewed Google sign-in"),
    ("A session starts", "I followed a session"),
    ("PasswordSecurity uses", "I reviewed PasswordSecurity, which uses"),
    ("The portal uses built-in roles", "I reviewed built-in roles"),
    ("Platform records connect", "I traced how platform records connect"),
    ("Announcements give", "I reviewed how announcements give"),
    ("The archive page gives users", "I used the archive page for"),
    ("Holiday records help", "I reviewed holiday records, which help"),
    ("Campus chat links", "I reviewed campus chat, which links"),
    ("A chat message moves", "I followed a chat message through"),
    ("ChatReminderService runs", "I reviewed ChatReminderService, which runs"),
    ("Profile pictures and platform images use", "I checked the upload routes used by"),
    ("The client uses static HTML", "I reviewed the client, which uses static HTML"),
    ("The administrator dashboard collects", "I reviewed the administrator dashboard, which collects"),
    ("The portal returns different responses", "I mapped the different responses"),
    ("Local checks showed", "During local checks, I saw"),
    ("A deployment starts", "I treated deployment as a sequence beginning with")
]


def humanize_text(text):
    for old, new in HUMANIZED_REWRITES:
        text = text.replace(old, new)
    return text.replace("’", "'").replace("“", '"').replace("”", '"')


def add_body(doc, text, after=7):
    paragraph = doc.add_paragraph(style="Normal")
    run = paragraph.add_run(humanize_text(text))
    set_run_font(run, size=10.5, color=BLACK)
    style_paragraph(paragraph, after=after)
    return paragraph


def add_small(doc, text, italic=False):
    paragraph = doc.add_paragraph(style="Normal")
    run = paragraph.add_run(text)
    set_run_font(run, size=9.2, color=MID_GRAY, italic=italic)
    style_paragraph(paragraph, after=5, line=1.05)
    return paragraph


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, line=1.0)
        r = p.add_run(header)
        set_run_font(r, size=9.2, color="FFFFFF", bold=True)
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    mark_header_row(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line=1.05)
            r = p.add_run(str(value))
            set_run_font(r, size=9.2, color=BLACK)
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            for i, width in enumerate(widths):
                cells[i].width = Inches(width)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    spacer = doc.add_paragraph()
    style_paragraph(spacer, after=3, line=1.0)
    return table


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run, size=10.3, color=BLACK)
    style_paragraph(paragraph, after=4, line=1.08)
    return paragraph


def add_figure(doc, image_path, caption, width=6.15):
    image_path = Path(image_path)
    if not image_path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    style_paragraph(paragraph, after=3, line=1.0)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run_font(run, size=9.2, color=MID_GRAY, italic=True)
    style_paragraph(cap, after=6, line=1.0)


def add_figure_pair(doc, left_path, left_caption, right_path, right_caption, width=3.05):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    mark_header_row(table.rows[0])
    for cell, image_path, caption in zip(table.rows[0].cells, [left_path, right_path], [left_caption, right_caption]):
        cell.width = Inches(width)
        set_cell_borders(cell, color="FFFFFF", size="0")
        set_cell_margins(cell, top=40, start=40, bottom=40, end=40)
        cell.text = ""
        image_path = Path(image_path)
        if image_path.exists():
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            shape = p.add_run().add_picture(str(image_path), width=Inches(width))
            shape._inline.docPr.set("descr", caption)
            shape._inline.docPr.set("title", caption)
            style_paragraph(p, after=3, line=1.0)
        cap = cell.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        set_run_font(run, size=8.5, color=MID_GRAY, italic=True)
        style_paragraph(cap, after=4, line=1.0)
    spacer = doc.add_paragraph()
    style_paragraph(spacer, after=2, line=1.0)


doc = Document()
doc.core_properties.title = "FIU Global Portal Internship Project Report"
doc.core_properties.subject = "Project review, implementation evidence, and internship learning record"
doc.core_properties.author = "[Your full name]"
doc.core_properties.keywords = "FIU Global Portal, internship, project report, web application"
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)
for style_name, size in (("Heading 1", 15), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLACK)

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.color.rgb = RGBColor.from_string(BLACK)
title_style_ppr = title_style._element.get_or_add_pPr()
title_style_borders = title_style_ppr.find(qn("w:pBdr"))
if title_style_borders is not None:
    title_style_ppr.remove(title_style_borders)

# Header and footer
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hr = hp.add_run("FIU GLOBAL PORTAL  |  INTERNSHIP PROJECT REPORT")
set_run_font(hr, size=8.5, color=BLACK, bold=True)
style_paragraph(hp, after=0, line=1.0)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("FINAL INTERNATIONAL UNIVERSITY  |  Page ")
set_run_font(fr, size=9, color=MID_GRAY)
add_page_field(fp)
style_paragraph(fp, after=0, line=1.0)

# Title page
title_space = doc.add_paragraph()
style_paragraph(title_space, after=10, line=1.0)
if LOGO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_shape = p.add_run().add_picture(str(LOGO), width=Inches(1.15))
    logo_shape._inline.docPr.set("descr", "FIU Global Portal logo")
    logo_shape._inline.docPr.set("title", "FIU Global Portal logo")
    style_paragraph(p, after=16, line=1.0)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
remove_paragraph_borders(title)
tr = title.add_run("FIU Global Portal")
set_run_font(tr, name="Aptos Display", size=28, color=BLACK, bold=True)
style_paragraph(title, after=6, line=1.0, keep=True)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = subtitle.add_run("Internship Project Report")
set_run_font(sr, name="Aptos", size=16, color=NAVY, bold=True)
style_paragraph(subtitle, after=24, line=1.0)

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
ir = intro.add_run("A practical review of the portal design, implementation, security, and test evidence")
set_run_font(ir, size=10.5, color=MID_GRAY, italic=True)
style_paragraph(intro, after=20, line=1.1)

details = [
    ("Prepared by", "[Your full name]"),
    ("Student number", "[Your student number]"),
    ("Department", "[Your department]"),
    ("Host organisation", "Final International University"),
    ("Internship period", "[Start date to end date]"),
    ("Supervisor", "[Supervisor name]"),
    ("Submission date", "[Date]")
]
table = doc.add_table(rows=0, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
for i, (label, value) in enumerate(details):
    cells = table.add_row().cells
    for cell in cells:
        set_cell_borders(cell)
        set_cell_margins(cell, top=120, bottom=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cells[0].width = Inches(1.75)
    cells[1].width = Inches(4.4)
    cells[0].text = ""
    cells[1].text = ""
    p0 = cells[0].paragraphs[0]
    p1 = cells[1].paragraphs[0]
    style_paragraph(p0, after=0, line=1.0)
    style_paragraph(p1, after=0, line=1.0)
    r0 = p0.add_run(label)
    r1 = p1.add_run(value)
    set_run_font(r0, size=9.8, color=NAVY, bold=True)
    set_run_font(r1, size=9.8, color=BLACK)
    set_cell_shading(cells[0], LIGHT_BLUE)
    set_cell_shading(cells[1], "FFFFFF")
mark_header_row(table.rows[0])
add_small(doc, "Version 1.0  |  Internship submission", italic=True)

doc.add_page_break()

# Opening content
add_heading(doc, "Executive summary", 1)
add_body(doc, "This report presents the FIU Global Portal, an internship project built to give students, instructors, and administrators one reliable entry point for university services. The portal brings service links, announcements, dining information, academic directory records, notifications, account settings, and campus chat into one sign-in flow.")
add_body(doc, "I reviewed the source structure, followed the main request paths, checked the local runtime, and recorded the build result. The project builds with zero warnings and zero errors. Local requests show the expected login redirect, a working login page, and protected data routes for users without a session. The next release needs automated journey tests, production monitoring, and a documented backup process.")
add_body(doc, "The report keeps starter records separate from live usage. The source includes 11 platform links, 4 LMS subplatforms, 9 faculties, and 33 departments in the supplied directory seed. These figures describe the project setup, not campus activity.")

add_heading(doc, "Contents", 1)
contents = [
    "1 Executive summary",
    "2 Project background and objectives",
    "3 Portal users and scope",
    "4 System design",
    "5 User portal features",
    "6 Administration features",
    "7 Security and privacy",
    "8 Accessibility and language support",
    "9 Internship work and learning",
    "9.1 Internship reflection",
    "10 Verification and results",
    "11 Constraints and risks",
    "12 Recommended next work",
    "13 Delivery assessment",
    "14 How to use the FIU Global Portal",
    "14.1 Portal screenshots, sign in and dashboard",
    "14.2 Portal screenshots, dining",
    "14.3 Portal screenshots, chat and administration",
    "15 Extended project record",
    "A Appendix Route summary"
]
for item in contents:
    add_bullet(doc, item)

add_heading(doc, "Project background and objectives", 1)
add_body(doc, "University services often run on separate sites. Users must remember several links and check each service for new information. This project addresses the access problem with one portal shell and a role-aware dashboard. A signed-in user sees assigned services, then opens the required destination from one screen.")
add_body(doc, "The project objectives were practical. Give users one trusted starting point. Show daily campus information in a short dashboard view. Give administrators direct control over content and access. Keep older accounts available during the migration from the earlier PHP setup. Protect account and chat data during normal use.")
doc.add_page_break()
add_table(doc, ["Project fact", "Observed value"], [
    ("Server", "ASP.NET Core on .NET 9 with minimal API routes"),
    ("Client", "Static HTML, CSS, and JavaScript pages"),
    ("Database", "MySQL through MySqlConnector"),
    ("User roles", "Student, instructor, administrator, super administrator, and custom roles"),
    ("Languages", "English, Turkish, French, Russian, and Arabic"),
    ("Starter directory", "9 faculties and 33 departments"),
    ("Build check", "Zero warnings and zero errors")
], widths=[1.75, 4.4])

add_heading(doc, "Portal users and scope", 1)
add_body(doc, "Students need fast access to academic tools, notices, meals, and personal account details. Instructors use the same daily services, with notification access and role options set by an administrator. Administrators maintain accounts and portal content. A super administrator also manages roles, permissions, and administrator records.")
add_body(doc, "The user portal covers dashboard shortcuts, platform links, announcements, dining menus, profile settings, notifications, archive records, and campus chat. The administration portal covers accounts, platforms, faculties, departments, announcements, dining menus, holidays, access rules, exports, and activity charts.")
add_body(doc, "The portal links to Leave and Absence, RMS, AIS, LMS, document forms, summer school, accommodation, support, exam registration, exemption forms, and resit applications. These links take users to the external university services. The portal stores the link details and role visibility rules.")

doc.add_page_break()
add_heading(doc, "System design", 1)
add_body(doc, "The server uses minimal API routes in Program.cs. AppDataStore contains the main portal operations. DatabaseStateStorage keeps the full state document in the dotnet_app_state table and mirrors core records into relational tables where older tools still need them. Sessions use dotnet_sessions. Chat messages use dotnet_chat_messages. Faculties and departments use linked directory tables.")
add_body(doc, "The application creates missing tables during startup. Middleware checks database access before protected pages and data routes. When MariaDB is offline, the host stays active and returns a retryable service response. You restore the database connection, then reload the page.")
add_body(doc, "The browser client loads page shells from wwwroot. JavaScript fetches data from the user and administrator API routes. A WebSocket route handles chat messages and delivery receipts. A hosted service checks pending reply reminders every 15 minutes and sends a message after 48 hours when SMTP settings exist.")
add_table(doc, ["Layer", "Main responsibility", "Source area"], [
    ("Presentation", "Login, dashboards, forms, calendars, dialogs, charts", "wwwroot HTML, CSS, and JavaScript"),
    ("API", "Authentication, role checks, data routes, uploads, exports", "Program.cs"),
    ("Domain state", "Accounts, roles, content, menus, holidays, activity", "Services/AppDataStore.cs"),
    ("Persistence", "JSON state, sessions, chat, academic directory", "Services/DatabaseStateStorage.cs"),
    ("Background work", "Pending chat reply reminders", "Services/ChatReminderService.cs")
], widths=[1.2, 3.0, 1.95])

add_heading(doc, "User portal features", 1)
add_body(doc, "The login page accepts local credentials and Google sign-in for final.edu.tr accounts. The server validates the Google state value, checks the returned identity, and rejects accounts outside the university domain. After sign-in, the dashboard shows most accessed services, today’s menu, and current announcements before the full platform list.")
add_body(doc, "Platform cards group services by section and role. The role access model supports section access and smaller permissions inside a section. This approach gives an administrator control over who sees a service and who gets permission to import or remove menu records.")
add_body(doc, "The dining page uses a month calendar. Each date opens breakfast and lunch details with meal times. Users with the required permission receive CSV or Excel import access. The announcements page uses a slider for current notices. The archive page keeps older notices and past menus available for later checks.")
add_body(doc, "The profile page supports name, faculty, department, profile picture, and password changes. Faculty selection filters the department list. Campus chat lets users search for people, open a conversation, send messages, and receive delivery and seen status updates.")

add_heading(doc, "Administration features", 1)
add_body(doc, "The administrator dashboard presents total users, active administrators, connected platforms, activity totals, role mix, and platform access records. Client-side charts show recent activity, account mix, platform access, and administrator actions.")
add_body(doc, "Account management covers user creation, edits, deletion, password changes, promotion to administrator, and demotion back to user. Administrators assign built-in roles or create a custom role. Platform management stores a section, name, description, destination URL, notification URL, image, and visible roles.")
add_body(doc, "The academic directory stores faculties and departments as linked records. Administrators add, edit, deactivate, search, and page through entries. Holiday management supports single records and bulk CSV or Excel import. Export and template downloads fit office workflows.")

add_heading(doc, "Security and privacy", 1)
add_body(doc, "Passwords use salted PBKDF2 with SHA-512. The current policy uses 210,000 iterations and a 32-byte derived value. Older SHA-256 or plain values receive a one-time upgrade after a successful sign-in. This keeps older accounts usable while moving them to the current password format.")
add_body(doc, "Session cookies use HttpOnly and SameSite protections. Secure cookies activate for HTTPS. Sessions expire after eight hours of inactivity. Expired records leave the active session store.")
add_body(doc, "Chat text uses AES-GCM encryption before storage. The key comes from environment settings. The server limits chat text to 2,000 characters and checks both users before saving a message. File uploads use a 10 MB request limit and accept common image formats.")
add_body(doc, "Role checks protect user and administrator routes. A request without a valid session receives status 401. A signed-in account without the needed permission receives status 403. Super administrator actions receive an extra role check. Google callback errors also clear the temporary state cookie.")

add_heading(doc, "Accessibility and language support", 1)
add_body(doc, "The pages include labels, live regions, keyboard-friendly buttons, alternative text, dialog roles, and status text for loading and error states. CSS media queries adjust the layout for phones, tablets, and larger screens. Reduced-motion preferences receive a dedicated style path. Dark mode appears in both portal shells.")
add_body(doc, "The portal includes English, Turkish, French, Russian, and Arabic. The language service checks the saved cookie, then the browser language, then English. Arabic uses right-to-left layout rules. User and administrator language choices persist with the account, so background messages use the saved preference.")

add_heading(doc, "Internship work and learning", 1)
add_body(doc, "The project gave me practice across the full web application path. I worked with C# route handlers, MySQL state storage, browser-side JavaScript, responsive CSS, form validation, file import logic, and access rules. I also followed how a login request becomes a session, how a role changes the visible dashboard, and how a user action reaches the database.")
add_body(doc, "The portal made security work concrete. I reviewed password migration, session cookie flags, Google domain checks, encrypted chat storage, upload limits, and permission checks. I also learned why a service needs a recovery response when a database connection fails. A clear retry path helps operations staff restore access without restarting the whole host.")
add_body(doc, "The main lesson was the value of small, testable flows. A platform card needs a valid target URL. A menu import needs row checks before a write. A chat message needs sender and recipient checks before storage. These small checks reduce faults at the point where a user action enters the system.")

doc.add_page_break()
add_heading(doc, "Internship reflection", 1)
add_body(doc, "At the start of the internship, I needed a reliable way to read a large project without losing the user view. I began with the page names, then followed each click into the route handlers and storage classes. This method gave me a simple record of what I saw and where each result came from.")
add_body(doc, "One test made the security model easy to understand. A request without a session returned status 401. The same boundary appeared on user and administrator routes. I wrote those results down before describing the role rules.")
add_body(doc, "I also learned to test a feature from the screen first. I opened the dining calendar, selected a dated menu, and captured the detail dialog. I used the same approach for chat, profile navigation, and the administrator dashboard. The screenshots in this report came from those checks.")
add_body(doc, "When I reviewed my own work, I noticed one gap. Build checks do not replace browser tests. My next step would be to automate the sign-in, menu, role, import, and chat paths before a live release. This sequence would give the team faster feedback on regressions.")

add_heading(doc, "Verification and results", 1)
add_body(doc, "I ran the project build with dotnet build --no-restore. The build completed with zero warnings and zero errors. I also checked the local host at 127.0.0.1:5099. The results below record only checks performed during this review.")
add_table(doc, ["Check", "Expected result", "Observed result"], [
    ("Root path", "Redirect to login", "302 redirect"),
    ("Login page", "HTML page loads", "200 response"),
    ("Session route without cookie", "Reject request", "401 response"),
    ("User data route without cookie", "Reject request", "401 response"),
    ("Administrator data route without cookie", "Reject request", "401 response"),
    ("Build", "Compile without errors", "0 warnings, 0 errors"),
    ("Application log", "Host starts without error output", "Listening on 127.0.0.1:5099")
], widths=[2.2, 2.2, 1.85])
add_body(doc, "Starter records in the source include 11 platform links, 4 LMS subplatforms, 2 sample users, 2 sample administrators, 2 announcements, 3 dining menus, 2 holidays, and 3 notifications. These records support local review. They do not measure live campus usage.")

add_heading(doc, "Constraints and risks", 1)
add_body(doc, "The repository has no dedicated automated test project. Build checks prove compilation, but they do not prove every login path, permission rule, import case, chat flow, or responsive layout. External university services also sit outside the portal code. A changed login page, broken target site, or expired certificate needs separate monitoring.")
add_body(doc, "Production use needs a secret store for database, Google, chat, and SMTP keys. The deployment also needs scheduled backups for the JSON state document and relational tables. The browser loads Google Fonts and Font Awesome from external hosts, so a strict network policy might change the visual result.")

doc.add_page_break()
add_heading(doc, "Recommended next work", 1)
add_table(doc, ["Priority", "Action", "Reason"], [
    ("1", "Add integration tests for login, Google callback errors, role access, profile updates, menu import, holiday import, and chat receipts.", "Covers the main server paths."),
    ("2", "Add browser tests for student, instructor, administrator, and custom-role journeys.", "Checks the page flow users see."),
    ("3", "Add login and chat rate limits, failed sign-in counts, and alerts.", "Reduces abuse risk and improves response."),
    ("4", "Set a backup schedule and test a restore before launch.", "Protects state and shortens recovery time."),
    ("5", "Monitor each external platform link and show service status.", "Makes external failures visible."),
    ("6", "Review image validation, file cleanup, and local font assets.", "Improves deployment control and storage hygiene.")
], widths=[0.7, 3.55, 2.0])

add_heading(doc, "Delivery assessment", 1)
add_body(doc, "FIU Global Portal has a broad working base for a university service hub. The source shows a complete user shell, a separate administrator shell, database persistence, role controls, multilingual content, responsive styling, secure password handling, encrypted chat storage, and practical import and export tools.")
add_body(doc, "The build and basic runtime checks pass. The next release needs automated journey tests, production monitoring, and a documented backup and recovery process. These steps will give you stronger evidence before a live campus launch.")

doc.add_page_break()
add_heading(doc, "How to use the FIU Global Portal", 1)
add_body(doc, "Use the portal in a short sequence. Start at the login page. Move through the dashboard. Open a service when needed. The steps below follow the screens captured during this review.")
for step in [
    "Open the portal address in your browser.",
    "Enter your username and password. Select Sign in. Use Google sign in for a final.edu.tr account when the option is enabled.",
    "Read the dashboard cards. Select Dashboard, Chat, Profile, Platforms, Announcements, or Dining menu from the left menu.",
    "Select Dining menu. Choose a date with a menu entry. Read breakfast and lunch details in the dialog.",
    "Select Chat. Search for a person, open the person row, type a message, and select Send. Review delivery and seen states.",
    "Select Profile. Update your name, faculty, department, picture, or password. Save the change, then review the status text.",
    "Use the administrator panel for record work. Review counts first. Open Users, Role Access, Manage Platforms, Faculties and Departments, Announcements, Dining Menu, or Holidays and Days Off. Save one change at a time.",
    "Sign out from the account menu after your work. A fresh sign in checks the session path."
]:
    add_bullet(doc, step)
add_body(doc, "Use the response status as a guide. Status 401 means sign in again. Status 403 means request the needed role. A retry message points to a database or service fault. Save the time and page name before reporting the fault.")

doc.add_page_break()
add_heading(doc, "Portal screenshots, sign in and dashboard", 1)
add_figure(doc, ROOT / "output" / "playwright" / "login.png", "Figure 1. Login page. Enter local credentials or choose Google sign in.")
add_figure(doc, ROOT / "output" / "playwright" / "student-dashboard.png", "Figure 2. Student dashboard. Review service cards, today's menu, and announcements.")

doc.add_page_break()
add_heading(doc, "Portal screenshots, dining", 1)
add_figure_pair(
    doc,
    ROOT / "output" / "playwright" / "student-dining.png",
    "Figure 3. Dining calendar. Choose a dated entry to view meal details.",
    ROOT / "output" / "playwright" / "student-dining-detail.png",
    "Figure 4. Dining detail dialog. Read breakfast and lunch times."
)

doc.add_page_break()
add_heading(doc, "Portal screenshots, chat and administration", 1)
add_figure_pair(
    doc,
    ROOT / "output" / "playwright" / "student-chat.png",
    "Figure 5. Student chat. Search for a person, open a conversation, and send a message.",
    ROOT / "output" / "playwright" / "admin-dashboard.png",
    "Figure 6. Administrator dashboard. Review counts and open management areas."
)

# Extended project record
extended_sections = [
    ("Internship setting and role", [
        "I treated the portal as a working product during the internship. My first task was to read the source before making judgments about the screens. I traced the server start-up path, the data store, the page scripts, and the protected routes. This gave me a practical view of how a small team keeps one product together.",
        "My work covered source reading, local running, route checks, build checks, and report writing. I recorded facts from the code rather than filling gaps with guesses. When a feature depended on an external service, I marked the boundary and kept the claim narrow.",
        "The role required steady attention to detail. A name in a menu, a role value, a cookie flag, or a table name links one part of the product to another. I learned to follow those links before changing a sentence or a test result."
    ], [
        "Read the server and client structure.",
        "Followed user and administrator request paths.",
        "Ran the local host and recorded responses.",
        "Mapped findings into a usable project report."
    ]),
    ("Project timeline", [
        "I split the work into four passes. The first pass covered orientation and source reading. The second pass covered runtime checks and feature mapping. The third pass covered security, access rules, and operational limits. The last pass covered report writing, layout, and review.",
        "This order helped me keep evidence close to each claim. I wrote down the route or source file first, then described the user effect. For example, a role check in a route became a short explanation of who receives access and who receives a rejection.",
        "The timeline also showed where future work belongs. Browser journey tests need a separate pass because a successful build does not prove a full page flow. Production monitoring and backup testing also need time outside local development."
    ], [
        "Pass one, source orientation.",
        "Pass two, local runtime checks.",
        "Pass three, security and operations review.",
        "Pass four, report preparation and visual review."
    ]),
    ("Requirements gathering", [
        "I gathered requirements from the existing page names, route handlers, models, seed records, and service classes. The user pages pointed to daily needs such as service links, menus, notices, profile settings, and chat. The administrator pages pointed to content control, account control, imports, exports, and role control.",
        "This approach kept the report close to the product. A requirement gained a place in the report when the source showed a page, a route, a model, or a stored record for the entry. A future idea stayed in the recommendation list instead of appearing as a completed feature.",
        "The result is a practical requirement set. Users need a short path from sign-in to a service. Staff need safe tools for changes. The portal needs clear responses when a database or external site is unavailable."
    ], [
        "User need, reach a university service quickly.",
        "Staff need, update records without direct database work.",
        "Security need, enforce role checks on every protected path.",
        "Operations need, recover from service and database faults."
    ]),
    ("Stakeholder map", [
        "The portal serves several groups with different duties. Students and instructors read daily information and open service links. Administrators maintain records and access. A super administrator manages the administrator set and role definitions. Operations staff support the database, mail settings, and deployment environment.",
        "I used this map when reviewing features. A student-facing action needs a short path and clear status text. An administrator action needs a permission check and a record of the change. An operations action needs a recovery path and a useful log message.",
        "The map also shows where handovers happen. The portal hands users to external university services. The browser hands form data to API routes. The server hands stored data to page scripts. Each handover needs a clear contract."
    ], None),
    ("User stories", [
        "User stories helped me read the portal from a daily point of view. A student wants to sign in, see the next meal, and open a learning service without searching through old bookmarks. An instructor wants the same quick view with access set by a role. An administrator wants to correct a menu record, change a platform link, or update a user without direct database edits.",
        "The stories also cover failure paths. A user without a session needs a clear return to login. A user without a permission needs a clear rejection. An administrator who submits a bad import needs a row-level message rather than a silent write.",
        "Writing these stories made the source easier to compare with the report. Each story points to a page, a route, a record type, or a permission rule."
    ], [
        "As a student, I want one dashboard for daily services.",
        "As an instructor, I want role-based service access.",
        "As an administrator, I want controlled content edits.",
        "As an operator, I want recovery evidence after a fault."
    ]),
    ("Functional requirements", [
        "The functional requirements describe actions a user or staff member performs. Sign-in creates a session. A dashboard request returns assigned platforms, current notices, and menu data. A profile update changes account details. A chat send creates a protected message and returns a delivery state.",
        "Administrator functions follow the same pattern. A platform edit validates the target fields before storage. A faculty edit keeps linked department records consistent. A menu import reads a file, checks row values, and stores accepted rows. A role edit changes the permission set used by later requests.",
        "These functions share a simple rule. Validate the request at the boundary, apply the business rule, store the result, and return a status for the page."
    ], [
        "Read, create, edit, deactivate, import, export, and search records.",
        "Return status text for success, rejection, and service failure.",
        "Keep user data and administrator data on separate route groups.",
        "Record activity for important administrator actions."
    ]),
    ("Nonfunctional requirements", [
        "The portal also has quality requirements. Pages need to load in a normal browser without a large client framework. Protected data needs a session and a role check. Passwords need a slow one-way derivation. Chat text needs encryption before database storage. Uploaded files need a size limit and type checks.",
        "The interface needs to work at phone, tablet, and desktop widths. Labels, status regions, dialog roles, and keyboard-friendly controls support users who rely on assistive tools. Language choices need to persist so users do not reset a preference on every visit.",
        "Operational quality matters as well. The host should keep running when the database drops, then return a retryable response. The deployment needs secret storage, backups, and a way to see failures outside the local machine."
    ], [
        "Security, privacy, and role control.",
        "Readable layouts and clear status messages.",
        "Stable responses during database faults.",
        "Simple deployment and repeatable recovery steps."
    ]),
    ("Data model", [
        "The domain model holds the records needed by a university service hub. Accounts carry identity, role, faculty, department, password data, language preference, and profile details. Platforms carry service labels, sections, target URLs, notification URLs, images, and visible roles. Notices, menus, holidays, faculties, and departments form the remaining core records.",
        "Links between records matter. A department points to a faculty. A platform has a set of visible roles. A chat record points to a sender and a recipient. A notification points to a user and a message state. These links give the portal enough context to return a focused response.",
        "The model also supports migration. Older account values stay readable during sign-in, then move to the current password format after a successful check. This avoids a forced reset for every account during the move from the earlier PHP setup."
    ], None),
    ("State document and relational tables", [
        "DatabaseStateStorage keeps a full state document in dotnet_app_state. The same service mirrors core records into relational tables where older tools still need them. This split supports the current portal while keeping links to earlier records and tools.",
        "The state document gives the application one place to load and save broad portal state. Relational tables give focused access to sessions, chat messages, faculties, and departments. The design also gives the start-up path a clear place to create missing tables.",
        "The trade-off needs documentation. A future change must update the state document rules and the mirrored tables in the same request. Backup work must include both forms. A restore test should confirm a user record, a menu record, a directory record, and a chat record after recovery."
    ], [
        "Full state, dotnet_app_state.",
        "Sessions, dotnet_sessions.",
        "Chat records, dotnet_chat_messages.",
        "Directory records, linked faculty and department tables."
    ]),
    ("API route design", [
        "Program.cs holds the minimal API route map. The route groups separate authentication, user data, administrator data, uploads, exports, and compatibility redirects. This keeps the request path visible in one server file while domain work stays in service classes.",
        "A protected route follows a small sequence. The handler reads the session, checks the role or permission, validates the request body, calls the store, and returns a status with data or an error message. The browser script reads the status and updates the page state.",
        "This design makes manual review easier. I requested a route without a cookie and recorded status 401. I requested an administrator route without a cookie and recorded the same boundary. Future integration tests should repeat these checks with valid and invalid roles."
    ], [
        "Authentication routes create and end sessions.",
        "User routes return personal and daily portal data.",
        "Administrator routes handle records and access rules.",
        "WebSocket routes handle live chat traffic."
    ]),
    ("Authentication flow", [
        "The local sign-in flow starts on login.html. The page sends credentials to the server. The server finds the account, checks the stored password value, creates a session, and sends a response with the next page or a clear error. A session cookie then travels with later requests.",
        "The root path checks the current session before deciding where to send the browser. An unauthenticated request receives a redirect to login. A signed-in user goes to the role-appropriate dashboard. The same boundary protects user data and administrator data routes.",
        "This flow gives the portal one place to enforce identity. Page scripts do not decide whether a user is trusted. The server repeats the check for every protected request."
    ], [
        "Read credentials.",
        "Check the stored password format.",
        "Create an eight-hour session record.",
        "Return a dashboard or an error status."
    ]),
    ("Google sign in checks", [
        "The portal also accepts Google sign-in for final.edu.tr accounts. The server creates a temporary state value, receives the callback, checks the state, and verifies the returned identity. The domain rule rejects accounts outside the university address space.",
        "The callback path checks more than an email string. The server validates the audience, the expiry, the returned identity, and the temporary state cookie. Failure paths clear the temporary cookie before returning an error. This reduces the chance of reusing an old callback state.",
        "Local review covered the route structure and rejection rules. A full provider test needs a configured client ID, secret, redirect URI, and a controlled test account. Those values do not belong in source control."
    ], [
        "State value check.",
        "Identity and audience check.",
        "Expiry check.",
        "final.edu.tr domain check.",
        "Temporary cookie cleanup on error."
    ]),
    ("Session lifecycle", [
        "A session starts after a successful local or Google sign-in. The server stores the session record in dotnet_sessions and places the session identifier in a protected cookie. Each protected request reads the record, checks its age, and refreshes activity when the account remains valid.",
        "Sessions expire after eight hours of inactivity. Expired records leave the active session store during cleanup. HttpOnly and SameSite flags limit browser exposure and cross-site use. Secure cookies activate when the host runs over HTTPS.",
        "The lifecycle gives operators a clear control point. A logout route removes the active record. A support worker who needs to end a session has one stored record to remove. Future monitoring should track session creation, expiry, and rejected requests without recording secret values."
    ], [
        "Create on successful sign-in.",
        "Check on every protected request.",
        "Expire after eight hours of inactivity.",
        "Remove on logout or cleanup."
    ]),
    ("Password security and migration", [
        "PasswordSecurity uses salted PBKDF2 with SHA-512. The current policy uses 210,000 iterations and a 32-byte derived value. A salt keeps equal passwords from producing equal stored values. The work factor makes repeated guessing more costly than a fast hash.",
        "The source still recognises older SHA-256 and plain values. After a successful sign-in, the server writes a new PBKDF2 value. This one-time upgrade keeps older accounts usable during migration and reduces the number of forced support resets.",
        "A production deployment needs a secret policy around database access and any administrative reset path. Support staff should never ask a user to send a password by email. A reset flow should issue a short-lived token and write an audit record."
    ], [
        "Generate a fresh salt per password.",
        "Use the current PBKDF2 policy for new values.",
        "Upgrade older values after a valid sign-in.",
        "Keep reset work outside normal password messages."
    ]),
    ("Authorization and custom roles", [
        "The portal uses built-in roles for student, instructor, administrator, and super administrator access. The model also supports custom roles with smaller permissions inside a service section. The server reads the role set before returning data or applying a change.",
        "A platform card appears for one role and stays hidden for another. A menu import needs a permission beyond normal page access. A super administrator action receives an extra role check because the action changes administrator records or role definitions.",
        "The model makes least privilege practical for a university setting. A staff member who edits dining data does not need full account control. A support account receives a narrow permission set and still completes a defined task."
    ], [
        "Check identity first.",
        "Check built-in or custom role next.",
        "Apply section and action permissions.",
        "Return 401 for no session and 403 for missing permission."
    ]),
    ("Platform link lifecycle", [
        "Platform records connect the portal to external university services. An administrator stores a section, name, description, target URL, notification URL, image, and visible roles. The user page groups cards by section and shows only records allowed for the current account.",
        "A link passes through several states during its life. Staff create a record, review the target, publish the record to selected roles, update the label or destination, and later hide or remove the record. Each step needs a clear success message and a way to correct a bad URL.",
        "External links sit outside the portal code. A link monitor should request each target on a schedule and show an outage status on the administrator page. A failed target should not block the rest of the dashboard."
    ], [
        "Create and edit link metadata.",
        "Choose visible roles.",
        "Review target and notification URLs.",
        "Monitor external response status."
    ]),
    ("Announcements and notifications", [
        "Announcements give the portal a daily information layer. The dashboard shows current notices before the full platform list. The announcements page uses a slider for current items, while the archive keeps older notices for later checks.",
        "Notifications use a user record and a message state. A user marks a message as read, and the page updates the visible count. An administrator creates, edits, publishes, or removes a notice according to the role rules.",
        "The useful design point is separation of current and older data. A busy dashboard stays short. A user who needs an earlier notice still has an archive path. Future work should add an expiry field and a clear publication date to reduce manual cleanup."
    ], [
        "Show current items on the dashboard.",
        "Keep older items in the archive.",
        "Track read state per user.",
        "Give staff a controlled publication flow."
    ]),
    ("Dining calendar and menu import", [
        "The dining page uses a month calendar. Each date opens breakfast and lunch details with meal times. This view fits a daily campus habit. Users scan a date, read the meal, and move on without opening a long list.",
        "Staff have two entry paths. They enter one menu record through a form, or they import rows from CSV or Excel. The import path needs permission checks, file checks, column checks, date checks, and a clear report of accepted and rejected rows.",
        "The archive keeps past menus available for later checks. A good operations routine exports a known month before a bulk edit and records who made the change. The source includes three starter dining menus for local review."
    ], [
        "Calendar view by month.",
        "Breakfast and lunch fields.",
        "CSV and Excel import paths.",
        "Archive access for earlier menus."
    ]),
    ("Archive and historical records", [
        "The archive page gives users a second view of older notices and menus. This matters when a student needs to check an earlier announcement or compare a past meal record. The archive keeps older content away from the short daily dashboard.",
        "Archive work needs a clear date rule. A notice becomes old after its active period ends. A menu remains useful after its date passes. The page should show the record date, title, and type so a user does not open the wrong item.",
        "Staff also need safe retention rules. A record should leave the archive only after a defined policy review. Export files should carry a date and owner so a later import does not erase the history by accident."
    ], [
        "Separate current and older records.",
        "Show dates and record types.",
        "Keep exports named and dated.",
        "Review retention rules with the university owner."
    ]),
    ("Faculty and department directory", [
        "The academic directory stores faculties and departments as linked records. The supplied seed contains 9 faculties and 33 departments. A user selects a faculty, then sees the matching department list. This keeps profile choices short and reduces spelling differences.",
        "Administrators add, edit, deactivate, search, and page through directory records. A faculty change needs care because linked departments depend on its identifier. A department should not become an orphan record after a faculty edit.",
        "The directory also supports future reporting. A platform or announcement might target a faculty or department once the permission model gains this scope. The current portal keeps the directory focused on profile data and searchable records."
    ], [
        "9 faculties in the supplied seed.",
        "33 departments in the supplied seed.",
        "Linked faculty and department identifiers.",
        "Search, paging, edit, and deactivate actions."
    ]),
    ("Holiday data management", [
        "Holiday records help the portal show dates when normal services change. Administrators add one holiday through a form or import a group from CSV or Excel. The import path fits calendar work at the start of a term or academic year.",
        "A reliable import checks the date format, title, duplicate dates, and empty rows before a write. A rejected row needs a line number and a reason. The accepted set should save in one controlled operation so a partial file does not leave an unclear calendar.",
        "Export and template downloads support office routines. Staff get a known column order before editing a file. An export also gives a small backup before a seasonal change."
    ], [
        "Single-record holiday form.",
        "Bulk CSV and Excel import.",
        "Template download for column order.",
        "Export for review and backup."
    ]),
    ("Chat and message privacy", [
        "Campus chat links users inside the portal. A user searches for a person, opens a conversation, writes a message, and sees status changes. ChatMessageProtector encrypts text with AES-GCM before storage. The encryption key comes from environment settings rather than page code.",
        "The server checks both users before saving a message. The server limits text to 2,000 characters and rejects a request when a sender or recipient is not valid. These checks keep a crafted request from writing to an unrelated conversation.",
        "Privacy also depends on operations. The key needs restricted storage and planned rotation. A rotation plan should record which messages use an old key and how support staff recover access during a controlled change."
    ], [
        "Encrypt before database storage.",
        "Check sender and recipient.",
        "Limit message length to 2,000 characters.",
        "Keep the encryption key outside source control."
    ]),
    ("Chat delivery receipts", [
        "A chat message moves through more than one state. The sender submits text. The server checks the conversation, stores the protected record, and sends a delivery event. The recipient page reports a seen state after the message enters the active view.",
        "The WebSocket route keeps the update live while the page stays open. A reconnect path remains important because a phone or campus network might drop the socket. The page should reload the last known messages before opening a new socket.",
        "Delivery and seen states also help support staff. A sender who reports a missing reply needs a time, a state, and a conversation record. The portal should avoid exposing message text in general logs."
    ], [
        "Send request.",
        "Store protected message.",
        "Send delivery receipt.",
        "Send seen update from the recipient view."
    ]),
    ("Reply reminder service", [
        "ChatReminderService runs as a hosted service. The service checks pending conversations every 15 minutes and sends a message after 48 hours when SMTP settings exist. The service supports a simple support habit, users receive a prompt when a reply waits too long.",
        "The reminder path needs idempotence. A check should mark a reminder as sent before the next interval finds the same conversation. A retry after a mail fault needs a safe state so one conversation does not receive repeated messages.",
        "SMTP settings belong in the deployment environment. Local review confirms the service class and interval. A production review should add a test mailbox, a mail failure alert, and a record of the last successful send."
    ], [
        "Check pending replies every 15 minutes.",
        "Use a 48-hour threshold.",
        "Send only when SMTP settings exist.",
        "Record send state and mail errors."
    ]),
    ("File upload and image handling", [
        "Profile pictures and platform images use upload routes. The server places a 10 MB request limit on uploads and accepts common image formats. The page shows a preview or a status message after the upload result returns.",
        "A safe upload path checks the file length, extension, content type, and decoded image data. The server should generate a storage name rather than trust a user file name. Old files need a cleanup rule when a profile picture changes.",
        "The current source provides a size limit and common format checks. A production hardening pass should add image dimension limits, malware scanning, storage quotas, and a rule for files left behind after a failed update."
    ], [
        "10 MB request limit.",
        "Common image format checks.",
        "Server-generated storage names.",
        "Cleanup for replaced and failed files."
    ]),
    ("Language service", [
        "The portal includes English, Turkish, French, Russian, and Arabic. The language service checks the saved cookie, then the browser language, then English. A saved account choice gives the user a stable setting across visits.",
        "Arabic uses right-to-left layout rules. The page needs more than translated labels. Direction changes affect menus, icon placement, text alignment, form order, and dialog buttons. The source keeps a dedicated RTL path for this reason.",
        "Background messages also follow the saved choice. This reduces mixed-language screens after a user changes a preference. Future content work should keep a translation key list and flag missing values before release."
    ], [
        "Five supported languages.",
        "Saved choice has first priority.",
        "Browser language has second priority.",
        "English acts as the fallback."
    ]),
    ("Responsive interface", [
        "The client uses static HTML, CSS, and JavaScript pages. CSS media queries adjust layouts for phones, tablets, and larger screens. The same page shell needs to hold a dashboard, a calendar, a table, or a dialog without hiding key actions.",
        "Responsive review should follow real tasks. Open the login page on a narrow screen. Read a menu date. Edit a profile field. Open a chat. Use an administrator table. A layout pass is useful only when the task still works at each width.",
        "The source also includes dark mode and reduced-motion rules. These options affect contrast, transition speed, focus visibility, and chart colours. A future browser test set should capture a few fixed viewport sizes and compare the main routes."
    ], [
        "Phone layout.",
        "Tablet layout.",
        "Desktop layout.",
        "Dark mode and reduced-motion checks."
    ]),
    ("Accessibility review", [
        "The pages include labels, live regions, keyboard-friendly buttons, alternative text, dialog roles, and status text for loading and error states. These details help users understand a page without relying on colour or pointer movement alone.",
        "I reviewed the source for labels and status messages, then ran an accessibility audit on the report file. The portal review remains a product task, so future work should run keyboard checks and screen-reader checks against every main route.",
        "A useful check follows a task from start to finish. A user should reach login fields, open a platform card, select a dining date, close a dialog, and read an error message with keyboard input. Focus should not disappear after a modal or route update."
    ], [
        "Label every form field.",
        "Announce loading and error state changes.",
        "Keep focus visible.",
        "Check keyboard and screen-reader paths."
    ]),
    ("Administrator dashboard metrics", [
        "The administrator dashboard collects working figures into one view. The dashboard shows total users, active administrators, connected platforms, activity totals, role mix, and platform access records. Client-side charts show recent activity and account mix.",
        "A metric helps staff answer a small question. How many accounts exist? Which platforms receive access? Which role group changed? Which administrator action happened most recently? The answer should link to a record list or an action history when staff need detail.",
        "The source includes starter data for local review. Those values support a screen check, not a claim about campus use. A production version needs a date range, a timezone rule, and a clear definition for each count."
    ], [
        "Total users.",
        "Active administrators.",
        "Platform access records.",
        "Recent activity and role mix."
    ]),
    ("Error handling and recovery", [
        "The portal returns different responses for different faults. A missing session receives 401. A valid session without permission receives 403. A database problem returns a retryable service response while the host stays active. A bad import returns a validation message instead of a blind write.",
        "This separation helps users and support staff. A user knows when login is needed. An administrator knows when access is missing. Operations staff know when the database needs repair. The page shows a short message while logs hold the deeper context.",
        "Recovery work should follow the same route after a fault. Restore the database connection, retry the request, and reload the page. A future test should force a database outage and confirm both the response and the recovery path."
    ], [
        "401, no valid session.",
        "403, missing permission.",
        "Retryable service response, database fault.",
        "Validation response, bad input."
    ]),
    ("Logging monitoring and operations", [
        "Local checks showed the host listening on 127.0.0.1:5099 with no error output in app-run.err. A production setup needs more than a start message. Operators need sign-in failure counts, route error counts, database health, mail failures, and external link status.",
        "Logs should answer who, what, when, and result without exposing passwords, chat text, or encryption keys. An administrator action needs an account identifier and an action name. A failed import needs a file name, row number, and reason.",
        "Monitoring should send an alert only after a useful threshold. One failed external link check might be a short outage. A repeated failure over several checks needs a service status on the administrator page and a support alert."
    ], [
        "Database health check.",
        "Sign-in failure count.",
        "Mail delivery failure count.",
        "External link status.",
        "Administrator action history."
    ]),
    ("Deployment checklist", [
        "A deployment starts with a clean build and a known configuration. The server needs the MySQL connection, Google settings, chat encryption key, and SMTP settings in the environment. The database user needs the smallest permission set for normal operation.",
        "Before launch, staff should apply the schema, load the approved directory and platform records, check the public login page, test a student account, test an administrator account, and confirm a rejected role request. The same steps should run after each release.",
        "A short checklist reduces missed work. The checklist gives a new operator a repeatable path and gives the project owner a place to sign off. The final check should record the version, date, operator, database state, and any known external service issue."
    ], [
        "Build with zero warnings and zero errors.",
        "Load secrets through environment settings.",
        "Run user and administrator smoke checks.",
        "Record release version and sign-off."
    ])
]

doc.add_page_break()
add_heading(doc, "Extended project record", 1)
add_body(doc, "The following pages record the internship work, design decisions, and operational checks in a consistent format.")
for section_index, (section_title, paragraphs, bullets) in enumerate(extended_sections, 1):
    if section_index > 1:
        doc.add_page_break()
    add_heading(doc, f"15.{section_index} {section_title}", 2)
    for paragraph_text in paragraphs:
        add_body(doc, paragraph_text)
    if bullets:
        for bullet_text in bullets:
            add_bullet(doc, bullet_text)

doc.add_page_break()
add_heading(doc, "Appendix Route summary", 1)
add_table(doc, ["Route", "Purpose", "Access"], [
    ("/login.html", "Local and Google sign-in page", "Public page"),
    ("/student_dashboard", "Student portal shell", "Authenticated user"),
    ("/instructor_dashboard", "Instructor portal shell", "Authenticated user"),
    ("/admin_dashboard", "Administrator portal shell", "Authenticated administrator"),
    ("/database/api.php", "User data and action routes", "Authenticated user"),
    ("/database/admin_api.php", "Administrator data and action routes", "Authenticated administrator"),
    ("/ws/chat", "Live chat connection", "Authenticated user"),
    ("/archive", "Archived notices and menus", "Authenticated user"),
    ("/rms_auth_bridge.php", "Compatibility redirect to RMS", "Public redirect")
], widths=[1.65, 3.05, 1.55])
add_small(doc, "Source files reviewed include Program.cs, Models/AppModels.cs, Services/AppDataStore.cs, Services/DatabaseStateStorage.cs, Services/PasswordSecurity.cs, Services/ChatMessageProtector.cs, the wwwroot portal pages, and the supplied academic directory seed script.", italic=True)

doc.save(OUTPUT)
print(OUTPUT)
